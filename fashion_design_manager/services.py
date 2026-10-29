"""Business services and transparent internal sustainability scoring."""
from datetime import datetime
from pathlib import Path
import re
from sqlalchemy import or_
from .database import Session, Design, Task, Product, Sustainability, Activity
from .config import EXPORT_DIR

def log_activity(message: str):
    with Session() as s: s.add(Activity(message=message)); s.commit()

def list_designs(query="", status="All"):
    with Session() as s:
        q=s.query(Design)
        if query:
            term=f"%{query}%"; q=q.filter(or_(Design.name.ilike(term),Design.category.ilike(term),Design.collection.ilike(term),Design.season.ilike(term),Design.sku.ilike(term),Design.designer.ilike(term)))
        if status != "All": q=q.filter(Design.status==status)
        return q.order_by(Design.updated_at.desc()).all()

def save_design(data, design_id=None):
    with Session() as s:
        d=s.get(Design,design_id) if design_id else Design()
        for k,v in data.items():
            if hasattr(d,k): setattr(d,k,v)
        s.add(d); s.commit(); s.refresh(d)
        s.add(Activity(message=f"{'Updated' if design_id else 'Created'} design · {d.name}")); s.commit(); return d.id

def delete_design(design_id):
    with Session() as s:
        d=s.get(Design,design_id)
        if d: s.delete(d); s.commit()

def sustainability_score(recycled=0, organic=0, recyclable=0, transport="", lifetime=1, water=0, energy=0):
    # Internal planning heuristic only; normalized values are explicitly not an LCA.
    material=min(100,(float(recycled)+float(organic))/2)
    recycle=min(100,float(recycled))
    packaging=min(100,float(recyclable))
    travel={"Rail":90,"Sea":75,"Road":45,"Air":10}.get(transport,50)
    life=min(100,float(lifetime)*20)
    impact=max(0,100-(float(water)/10+float(energy)))
    return round(material*.25+recycle*.20+packaging*.15+travel*.15+life*.15+impact*.10)

def save_sustainability(design_id, values):
    score=sustainability_score(**{k:values.get(k,0) for k in ("recycled","organic","recyclable","transport","lifetime","water","energy")})
    with Session() as s:
        row=s.query(Sustainability).filter_by(design_id=design_id).first()
        if row is None: row=Sustainability(design_id=design_id)
        for k,v in values.items(): setattr(row,k,v)
        row.score=score; d=s.get(Design,design_id); d.score=score; s.add_all([row,d]); s.commit()
    log_activity(f"Updated sustainability profile · design #{design_id}"); return score

def create_task(title, assignee="Studio team", due_date="", design_id=None):
    with Session() as s:
        t=Task(title=title,assignee=assignee,due_date=due_date,design_id=design_id); s.add(t); s.commit(); s.refresh(t); s.add(Activity(message=f"Created task · {title}")); s.commit(); return t.id

def list_tasks():
    with Session() as s: return s.query(Task).order_by(Task.id.desc()).all()

def dashboard_data():
    with Session() as s:
        ds=s.query(Design).all(); tasks=s.query(Task).all(); products=s.query(Product).all()
        return {"designs":len(ds),"analyzed":sum(bool(d.prediction) for d in ds),"sustainable":sum(d.score>=70 for d in ds),"tasks":sum(t.status!="Done" for t in tasks),"products":len(products),"categories":_counts([d.category for d in ds]),"statuses":_counts([d.status for d in ds]),"activities":s.query(Activity).order_by(Activity.created_at.desc()).limit(8).all(),"tasks_all":tasks,"avg_score":round(sum(d.score for d in ds)/max(1,len(ds)))}

def _counts(values):
    out={}
    for v in values: out[v or "Unassigned"]=out.get(v or "Unassigned",0)+1
    return out

def save_product(design_id,title,sku,description,price,inventory,platform_id=""):
    with Session() as s:
        p=Product(design_id=design_id,title=title,sku=sku,description=description,price=price,inventory=inventory,platform_id=platform_id); s.add(p); s.commit(); s.add(Activity(message=f"Prepared product listing · {title}")); s.commit(); return p.id

def export_excel(kind="catalog", path=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from .database import Session, Design, Sustainability, Task
    path=Path(path) if path else EXPORT_DIR / f"{kind}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    path.parent.mkdir(parents=True,exist_ok=True); wb=Workbook(); ws=wb.active
    if kind=="sustainability":
        ws.title="Sustainability"; headers=["Design","Material","Source","Recycled %","Organic %","Packaging","Recyclable %","Transport","Lifetime (years)","Internal Score"]
        with Session() as s: rows=[(x.design.name,x.design.material,x.source,x.recycled,x.organic,x.packaging,x.recyclable,x.transport,x.lifetime,x.score) for x in s.query(Sustainability).all()]
    elif kind=="production":
        ws.title="Production"; headers=["SKU","Product","Quantity","Status","Deadline","Responsible"]
        with Session() as s: rows=[(t.design.sku if t.design else "",t.title,1,t.status,t.due_date,t.assignee) for t in s.query(Task).all()]
    else:
        ws.title="Design Catalog"; headers=["Design ID","Design Name","Category","Collection","Season","Material","Color","AI Prediction","AI Confidence","Sustainability Score","Status","Designer","Created Date"]
        with Session() as s: rows=[(d.id,d.name,d.category,d.collection,d.season,d.material,d.color,d.prediction,d.confidence,d.score,d.status,d.designer,d.created_at.strftime("%Y-%m-%d")) for d in s.query(Design).all()]
    ws.append(headers)
    for row in rows: ws.append(row)
    for c in ws[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="23463F"); c.alignment=Alignment(vertical="center")
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
    for col in ws.columns:
        width=min(36,max(13,max(len(str(cell.value or "")) for cell in col)+2)); ws.column_dimensions[get_column_letter(col[0].column)].width=width
    wb.save(path); log_activity(f"Exported {kind} workbook"); return path

def safe_asset_name(name): return re.sub(r"[^A-Za-z0-9_-]+","_",name).strip("_")[:80] or "design"
